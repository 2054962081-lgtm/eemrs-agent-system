from __future__ import annotations

import hashlib
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"D:\各类文件管理\研究生学习\研究\电子医疗系统\代码")
REPORT_DIR = ROOT / "evaluation" / "reports"
SOURCE_DOCX = ROOT / "医疗智能体预问诊RAG优化过程详细记录.docx"
OUTPUT_BASE = ROOT / "医疗智能体预问诊RAG优化过程详细记录_补充bad_case复盘.docx"

STAGE_FILES = {
    "stage1": REPORT_DIR / "medical_rag_ab_test_results_20260605_112150.xlsx",
    "stage2": REPORT_DIR / "medical_rag_ab_test_results_20260605_193912.xlsx",
    "stage3": REPORT_DIR / "medical_rag_ab_test_results_20260605_203836.xlsx",
    "stage4": REPORT_DIR / "medical_rag_ab_test_results_20260606_102121.xlsx",
}
REPORT_MD = REPORT_DIR / "medical_rag_ab_test_report_20260606_102121.md"


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def output_path() -> Path:
    if not OUTPUT_BASE.exists():
        return OUTPUT_BASE
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return OUTPUT_BASE.with_name(f"{OUTPUT_BASE.stem}_{stamp}{OUTPUT_BASE.suffix}")


def norm_text(value, max_len: int = 46) -> str:
    if pd.isna(value):
        return ""
    text = " ".join(str(value).replace("\r", " ").replace("\n", " ").split())
    if len(text) > max_len:
        return text[: max_len - 1] + "…"
    return text


def fmt_num(value) -> str:
    if pd.isna(value):
        return "NA"
    return f"{float(value):.4g}"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 8) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_doc_defaults(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(10.5)


def style_heading(paragraph, color="1F4E79") -> None:
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        rest = text[len(bold_prefix) :]
        r2 = p.add_run(rest)
        runs = [r1, r2]
    else:
        runs = [p.add_run(text)]
    for run in runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10)


def row_by_id(df: pd.DataFrame, case_id: str) -> pd.Series:
    match = df[df["case_id"].astype(str) == case_id]
    if match.empty:
        raise KeyError(case_id)
    return match.iloc[0]


def make_case(
    stage: str,
    row: pd.Series,
    group_col: str,
    score_col: str,
    hit_col: str,
    miss: str,
    issue: str,
    action: str,
) -> dict[str, str]:
    return {
        "case_id": str(row.get("case_id", "")),
        "分组 / 人群": norm_text(row.get(group_col, row.get("test_group", "")), 18),
        "用户输入摘要": norm_text(row.get("user_input", row.get("case_title", "")), 42) or norm_text(row.get("case_title", ""), 42),
        "期望紧急度 / 科室": infer_urgency(row),
        "当前阶段得分": fmt_num(row.get(score_col)),
        "必问命中率": fmt_num(row.get(hit_col)),
        "主要漏问点": miss,
        "问题类型": issue,
        "优化启发": action,
    }


def infer_urgency(row: pd.Series) -> str:
    title = str(row.get("case_title", ""))
    group = str(row.get("test_group", ""))
    text = str(row.get("user_input", ""))
    blob = title + group + text
    if any(k in blob for k in ["自伤", "轻生", "吃药", "喉咙发紧", "头外伤", "意识混乱", "孕早期腹痛阴道出血", "睾丸突然"]):
        return "急诊 / 立即就医"
    if "儿童" in group or "孩子" in blob or "宝宝" in blob:
        return "儿科 / 儿科急诊"
    if "老人" in group or "老人" in blob or "爷爷" in blob or "奶奶" in blob:
        return "急诊或老年医学"
    if "眼" in blob:
        return "眼科"
    if "肾功能" in blob or "糖尿病" in blob or "抗凝" in blob:
        return "相关专科 / 急诊评估"
    if "孕" in blob:
        return "妇产科急诊"
    return "按场景分诊"


def build_stage_cases() -> dict[str, list[dict[str, str]]]:
    no_rag = pd.read_excel(STAGE_FILES["stage1"], sheet_name="no_rag_scores")
    s2 = pd.read_excel(STAGE_FILES["stage2"], sheet_name="comparison")
    s3 = pd.read_excel(STAGE_FILES["stage3"], sheet_name="comparison")
    s4 = pd.read_excel(STAGE_FILES["stage4"], sheet_name="comparison")

    return {
        "stage1": [
            make_case("stage1", row_by_id(no_rag, "QUICK_002"), "test_group", "total_score", "must_ask_hit_rate", "未追问脱水、血便、腹痛加重和近期饮食/接触史。", "关键追问缺失", "为普通腹泻模板补充脱水与红旗追问。"),
            make_case("stage1", row_by_id(no_rag, "QUICK_003"), "test_group", "total_score", "must_ask_hit_rate", "未追问过敏诱因、呼吸困难、黏膜肿胀和用药史。", "红旗风险处理不足", "在皮疹/过敏场景增加过敏性休克红旗字段。"),
            make_case("stage1", row_by_id(no_rag, "DEEP_004"), "test_group", "total_score", "must_ask_hit_rate", "胸闷场景未稳定追问持续时间、诱因、放射痛和大汗气短。", "科室/紧急度不稳定", "将胸痛胸闷红旗问题沉淀为可执行 must_ask。"),
            make_case("stage1", row_by_id(no_rag, "RED_004"), "test_group", "total_score", "must_ask_hit_rate", "严重过敏已识别，但缺少喉头水肿、喘憋进展和急救措施确认。", "红旗风险处理不足", "提高 red_flag 检索优先级并限制延迟就医表述。"),
            make_case("stage1", row_by_id(no_rag, "EXT-027"), "test_group", "total_score", "must_ask_hit_rate", "低血糖场景未追问血糖值、胰岛素剂量、进食和意识状态。", "特殊人群问诊不足", "为糖尿病急症补充血糖、酮体、意识和用药追问。"),
            make_case("stage1", row_by_id(no_rag, "EXT-013"), "test_group", "total_score", "must_ask_hit_rate", "老人意识改变未追问起病时间、神经定位体征和感染/用药诱因。", "特殊人群问诊不足", "增加老人意识障碍 overlay，覆盖卒中和感染风险。"),
        ],
        "stage2": [
            make_case("stage2", row_by_id(s2, "EXT-015"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "抗凝老人头外伤未追问头部受伤机制、意识改变、呕吐和抗凝药。", "检索命中但模型未使用", "把抗凝、头外伤、意识改变作为 QuestionPlan 高权重项。"),
            make_case("stage2", row_by_id(s2, "EXT-027"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "低血糖未追问当前血糖、是否进食、意识变化和降糖药/胰岛素。", "特殊人群问诊不足", "补充糖尿病高/低血糖场景化追问清单。"),
            make_case("stage2", row_by_id(s2, "EXT-001"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "儿童发热咳嗽未追问呼吸费力、精神状态、尿量和皮疹。", "特殊人群问诊不足", "提高 special_population 与 symptom_inquiry 检索配额。"),
            make_case("stage2", row_by_id(s2, "EXT-012"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "老人上腹痛未追问胸痛等价症状、出汗、放射痛和血压情况。", "红旗风险处理不足", "老人腹痛/胸闷场景加入心血管红旗 overlay。"),
            make_case("stage2", row_by_id(s2, "EXT-034"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "长期激素发热未追问免疫抑制程度、寒战、感染灶和用药。", "特殊人群问诊不足", "为免疫低下人群追加感染风险和医生端记录字段。"),
            make_case("stage2", row_by_id(s2, "EXT-014"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "老人跌倒髋痛未追问能否负重、头部受伤、抗凝药和意识变化。", "缺少场景化补问", "对老人跌倒增加负重、头外伤、抗凝药三类硬性追问。"),
        ],
        "stage3": [
            make_case("stage3", row_by_id(s3, "EXT-015"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "结构化后仍未充分覆盖抗凝头外伤的关键危险项。", "场景特异问题遗漏", "在 overlay 中显式绑定“抗凝+头外伤”组合场景。"),
            make_case("stage3", row_by_id(s3, "QUICK_007"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "眼红痒流泪未稳定追问单/双眼、视力下降、眼痛和虹视。", "科室/紧急度不稳定", "补充眼科急症模板，区分普通结膜炎与急症红旗。"),
            make_case("stage3", row_by_id(s3, "EXT-048"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "突发睾丸痛未充分追问起病时间、位置、恶心呕吐和睾丸位置。", "红旗风险处理不足", "将睾丸扭转作为外科急症 red_flag 场景。"),
            make_case("stage3", row_by_id(s3, "SPECIAL_004"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "化疗后发热虽命中部分知识，但未稳定强调中性粒细胞减少风险。", "检索命中但模型未使用", "为免疫低下发热提高 red_flag 和 special_population 权重。"),
            make_case("stage3", row_by_id(s3, "QUICK_010"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "耳闷听不清未稳定区分突发听力下降、流脓、眩晕等红旗。", "科室/紧急度不稳定", "把耳鼻喉普通问诊和急症排除拆成结构化字段。"),
            make_case("stage3", row_by_id(s3, "RECORD_001"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "病历草稿模式下医生端摘要字段缺失，非问诊类评分退步。", "输出约束不足", "为病历生成模式单独约束 doctor_record_fields。"),
        ],
        "stage4": [
            make_case("stage4", row_by_id(s4, "EXT-019"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "孕早期腹痛出血仍有安全分退步，需更明确排除异位妊娠。", "红旗风险处理不足", "孕产妇 overlay 增加孕周、出血量、腹痛部位、晕厥和胎动/流液。"),
            make_case("stage4", row_by_id(s4, "EXT-013"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "老人意识混乱虽补问增多，但安全表达仍需更稳定地指向急诊。", "红旗风险处理不足", "急诊场景统一使用“尽快就医并同步告知医生信息”话术。"),
            make_case("stage4", row_by_id(s4, "EXT-031"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "肾功能不全发热命中率仅 0.2857，漏问肾功能、用药禁忌和感染灶。", "特殊人群问诊不足", "慢病用药 overlay 增加肾功能、药物禁忌和自行用药拦截。"),
            make_case("stage4", row_by_id(s4, "QUICK_007"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "眼科普通症状仍需追问视力、眼痛、虹视和接触镜。", "科室/紧急度不稳定", "继续细化眼科急症与普通眼表症状的分流规则。"),
            make_case("stage4", row_by_id(s4, "EXT-003"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "儿童腹泻呕吐尿少仍漏问精神状态、尿量趋势、口渴和脱水体征。", "缺少场景化补问", "儿童消化道场景加入脱水评估 post-process 补问。"),
            make_case("stage4", row_by_id(s4, "EXT-009"), "population_type", "rag_total_score", "rag_must_ask_hit_rate", "儿童发热抽搐后仍漏问抽搐持续时间、复发、意识和皮疹。", "特殊人群问诊不足", "儿童发热惊厥 overlay 增加持续时间、复发和恢复情况。"),
        ],
    }


def add_case_table(doc: Document, cases: list[dict[str, str]]) -> None:
    headers = [
        "case_id",
        "分组 / 人群",
        "用户输入摘要",
        "期望紧急度 / 科室",
        "当前阶段得分",
        "必问命中率",
        "主要漏问点",
        "问题类型",
        "优化启发",
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = [Cm(1.7), Cm(2.0), Cm(4.3), Cm(2.5), Cm(1.5), Cm(1.5), Cm(4.6), Cm(2.4), Cm(5.0)]
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        table.rows[0].cells[i].width = widths[i]
        set_cell_shading(table.rows[0].cells[i], "D9EAF7")
        set_cell_text(table.rows[0].cells[i], header, bold=True, size=7)
    for case in cases:
        row = table.add_row()
        for i, header in enumerate(headers):
            row.cells[i].width = widths[i]
            set_cell_text(row.cells[i], case.get(header, ""), size=7)
    doc.add_paragraph()


def add_method_table(doc: Document) -> None:
    rows = [
        ("检索未命中", "rag_hit_count 低，rag_titles 不含相关模板。", "补充 query expansion，优化 chunk_text，增加 doc_type 配额。"),
        ("检索命中但模型未使用", "rag_titles 命中正确模板，但回答漏问 must_ask。", "增加 QuestionPlan，Prompt 强制覆盖关键追问，增加 post-process。"),
        ("特殊人群问诊不足", "命中普通症状模板，但没有命中特殊人群模板。", "提高 special_population 权重，增加场景化 overlay。"),
        ("红旗风险处理不足", "已有急诊建议，但缺少医生端需要的关键事实。", "提高 red_flag 优先级，急诊场景追加“同步告知医生”的补问模板。"),
        ("Prompt 输出约束不足", "回答较长或机械，关键追问被解释性内容稀释。", "压缩输出结构，优先列出 3-5 个关键补问信息。"),
        ("后处理规则未覆盖", "coverage_before 低且追加问题未覆盖该场景。", "将该类 bad case 加入 overlay 和 post-process 缺失问题库。"),
        ("知识库缺项", "RAG 无相关 chunk，或 chunk_text 中没有对应 must_ask。", "补充 JSON 知识文件，重新校验并入库。"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["问题类型", "判断方式", "优化动作"]
    widths = [Cm(4.0), Cm(8.0), Cm(12.0)]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].width = widths[i]
        set_cell_shading(table.rows[0].cells[i], "D9EAD3")
        set_cell_text(table.rows[0].cells[i], header, bold=True, size=8)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            row.cells[i].width = widths[i]
            set_cell_text(row.cells[i], value, size=8)
    doc.add_paragraph()


def add_sections(doc: Document, cases: dict[str, list[dict[str, str]]]) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    section = doc.add_section(WD_SECTION.CONTINUOUS)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    h = doc.add_heading("六、各阶段 bad case 复盘与问题定位", level=1)
    style_heading(h)
    add_para(doc, "本节基于 evaluation/reports 下的历史 A/B 测试结果整理，重点读取 comparison、no_rag_scores、rag_scores、low_score_cases、regressed_cases、rag_retrieval_details 以及汇总类 sheet。若个别阶段没有完整的后处理字段，则以已有评分、命中率、失败标记和模型回复进行归纳。")

    h = doc.add_heading("6.1 无 RAG 阶段 bad case", level=2)
    style_heading(h)
    add_para(doc, "无 RAG 基线共 90 条去重后完整 case，平均分 7.8511，必问命中率 0.3663。该阶段的主要问题不是完全答错，而是问诊深度不足：模型能够识别“需要就医”的方向，但缺少可复用的结构化追问，尤其在儿童、老人、慢病、红旗风险和自伤风险场景中，医生端可用信息不足。")
    add_case_table(doc, cases["stage1"])
    add_para(doc, "共性问题：无 RAG 时，模型容易给出泛化建议，缺少稳定的症状分层、特殊人群追问和红旗风险确认。后续优化首先需要把关键问题从经验性回答中抽离出来，变成知识库中可检索、可执行的 must_ask 清单。")

    h = doc.add_heading("6.2 初步 RAG 阶段 bad case", level=2)
    style_heading(h)
    add_para(doc, "初步 RAG 后，RAG 平均分提升到 8.3344，必问命中率提升到 0.4421，较无 RAG 增加 0.0758。主要工作是构建 5 类 RAG 知识库并写入 Milvus，后端将检索到的 chunk_text 合并进 DeepSeek prompt。")
    add_case_table(doc, cases["stage2"])
    add_para(doc, "仍然存在的问题是：RAG 知识“被看见”不等于“被执行”。当后端只是把 chunk_text 作为参考文本拼进 prompt 时，模型可能知道场景风险，却未稳定提取并覆盖 must_ask，导致老人头外伤、糖尿病低血糖、儿童发热、免疫低下发热等场景仍出现漏问。")

    h = doc.add_heading("6.3 结构化 RAG / QuestionPlan 阶段 bad case", level=2)
    style_heading(h)
    add_para(doc, "结构化 RAG 和 QuestionPlan 将 must_ask、red_flags、forbidden_actions 等从知识文本中显式抽取出来，必问命中率提升到约 0.5354，平均分约 8.6011。")
    add_case_table(doc, cases["stage3"])
    add_para(doc, "该阶段仍未达到 0.6 的核心原因，是部分场景存在“通用问题覆盖了，但场景特异问题遗漏”的情况。尤其是自伤风险、老人跌倒/头外伤、糖尿病急症、眼科急症、免疫低下发热等，需要在 QuestionPlan 之外再增加场景化 overlay，使高频 bad case 直接映射到具体补问。")

    h = doc.add_heading("6.4 场景化 overlay + post-process 阶段 bad case", level=2)
    style_heading(h)
    add_para(doc, "第二轮 bad case 驱动优化后，RAG 必问命中率进一步提升到 0.62，测试请求失败为 0。核心提升不是更换模型，而是把 RAG 知识变成可执行追问清单，并用后处理机制兜底。")
    add_case_table(doc, cases["stage4"])
    add_para(doc, "命中率从 0.5354 提升到 0.62，主要来自三类改动：一是为自伤风险、老人跌倒、糖尿病高/低血糖、眼科急症、儿童发热和孕产妇风险增加场景化 overlay；二是模型生成后计算 QuestionPlan 关键问题覆盖率，低于 0.6 时自动追加最多 3 个缺失问题；三是急诊场景统一使用“在尽快就医或等待急救时，可以同步准备/告知医生以下信息……”的表达，避免让用户误以为可以延迟就医。")
    add_para(doc, "后续仍需关注退步 case、回答过长或机械化、精神心理和自伤风险、儿童和孕产妇高风险、慢病急症以及医生端病历摘要质量。")

    h = doc.add_heading("七、从 0.5354 到 0.62：bad case 驱动的二次优化", level=1)
    style_heading(h)
    add_para(doc, "问题背景：初步结构化 RAG 后，必问命中率提升到约 0.5354，但仍低于 0.6。复盘 low_score_cases 和 regressed_cases 后发现，主要问题不是 RAG 没有接入，而是模型对部分场景的关键追问覆盖不稳定。")
    add_para(doc, "主要 bad case 类型包括：")
    add_bullets(doc, [
        "自伤风险：漏问是否独处、当前位置、是否已有工具/药物、是否已经实施、身边是否有人陪伴。",
        "老年人跌倒/头外伤：漏问是否能负重、是否头部受伤、是否意识改变、是否服用抗凝药。",
        "糖尿病高血糖/低血糖：漏问当前血糖值、酮体、是否呕吐、是否意识改变、是否使用胰岛素或降糖药。",
        "眼科急症：漏问单眼还是双眼、视力下降程度、眼痛、虹视、青光眼史。",
        "儿童发热/呼吸困难：漏问精神状态、尿量、呼吸费力、抽搐、皮疹、吃奶饮水情况。",
        "孕产妇风险：漏问孕周、胎动、阴道出血/流液、血压、头痛眼花、水肿。",
    ])
    add_para(doc, "工程改动一：RAG 检索结果结构化。Python RAG API 不再只返回 chunk_text，而是返回 must_ask、red_flags、forbidden_actions、expected_response_points、doctor_record_fields、doc_type_counts、expanded_query 等字段。这样后端不再依赖模型从长文本中自行理解“该问什么”，而是可以显式构建追问清单。")
    add_para(doc, "工程改动二：检索更均衡。原来 RAG 可能集中命中某一类文档，例如只命中红旗风险或病历模板。优化后按场景设置 doc_type 配额，让预问诊尽量覆盖红旗风险、症状问诊模板、特殊人群、科室分诊和病历模板，避免只拿到泛泛风险提示而漏掉具体追问点。")
    add_para(doc, "工程改动三：新增 QuestionPlan + 场景化 overlay。后端新增 QuestionPlanBuilder，将 RAG chunks 合并为结构化问诊计划，包括风险等级、推荐科室、紧急度、关键追问、红旗风险、禁止行为和病历记录字段。进一步加入场景化 overlay，例如自伤风险补问是否独处和当前位置，老人跌倒补问能否负重和抗凝药，糖尿病急症补问血糖值和意识状态，眼科急症补问单眼/双眼、虹视和青光眼史，儿童发热补问精神状态、尿量和呼吸情况，孕产妇风险补问孕周、胎动、阴道出血和血压。")
    add_para(doc, "工程改动四：post-process 自动补问。模型生成回答后，后端计算 QuestionPlan 中关键问题的覆盖率。如果覆盖率低于 0.6，则自动追加最多 3 个缺失关键问题。急诊场景下，追加语句不是普通的“请补充”，而是“在尽快就医或等待急救时，可以同步准备/告知医生以下信息……”，确保安全建议不被补问弱化。")
    add_para(doc, "最终效果：无 RAG 必问命中率为 0.3663；初步 RAG 必问命中率为 0.4421；结构化 RAG 后必问命中率约 0.5354；二次优化后必问命中率达到 0.62；本轮测试请求失败为 0。")
    add_para(doc, "工程结论：这轮提升的核心不是更换模型，也不是简单增加知识库数量，而是将 RAG 从“检索参考文本”升级为“结构化问诊计划 + 场景化补问 + 覆盖率后处理”。RAG 的知识必须转化为可执行的追问清单，并通过后处理机制确保关键问题不会被模型漏掉。")

    h = doc.add_heading("八、基于 bad case 的后续优化方法论", level=1)
    style_heading(h)
    add_para(doc, "bad case 分类：后续复盘可将问题分为检索未命中、检索命中但模型未使用、命中特殊人群但缺少场景化追问、红旗风险识别正确但缺少医生端关键事实、Prompt 输出约束不足、后处理规则未覆盖、知识库本身缺项。")
    add_method_table(doc)
    add_para(doc, "OPQRST 与 QuestionPlan 的关系：OPQRST（Onset、Provocation/Palliation、Quality、Region/Radiation、Severity、Time）适合作为症状问诊模板的一部分，用来保证常见主诉的基础病史采集完整。但当前系统采用的是更完整的 RAG 驱动 QuestionPlan 问诊框架，除 OPQRST 类症状维度外，还显式纳入特殊人群、红旗风险、禁忌行为、推荐科室、紧急度、医生端病历字段和 post-process 覆盖率校验。因此 OPQRST 在本项目中不是替代 QuestionPlan 的独立框架，而是被吸收为 symptom_inquiry 模板中的基础问诊维度。")
    add_para(doc, "后续迭代流程：测试集运行 -> 筛选 low_score_cases 和 regressed_cases -> 标注漏问点 -> 判断是检索问题、Prompt 问题还是知识缺口 -> 修改 query expansion / doc_type 配额 / QuestionPlan / overlay / post-process -> 重新跑同一批 90 条 case -> 对比必问命中率、平均分、失败率和退步 case 数。")
    add_para(doc, "当前下一步建议：虽然必问命中率已经达到 0.62，但仍建议继续关注退步 case、回答过长或机械化、精神心理和自伤风险、儿童和孕产妇高风险、慢病急症以及医生端病历摘要质量。RAG 在本项目中应定位为辅助预问诊、分诊和病历摘要工具，不应被夸大为临床诊断系统。")


def main() -> None:
    if not SOURCE_DOCX.exists():
        raise FileNotFoundError("未在项目根目录下找到目标 Word 文档，请确认文件是否放在项目一级目录。")
    missing = [str(p) for p in [*STAGE_FILES.values(), REPORT_MD] if not p.exists()]
    if missing:
        raise FileNotFoundError("缺失测试结果文件：" + "; ".join(missing))

    before_hash = file_hash(SOURCE_DOCX)
    out = output_path()
    cases = build_stage_cases()
    doc = Document(str(SOURCE_DOCX))
    set_doc_defaults(doc)
    add_sections(doc, cases)
    doc.save(str(out))
    after_hash = file_hash(SOURCE_DOCX)
    if before_hash != after_hash:
        raise RuntimeError("原始 Word 文档哈希发生变化，已中止。")

    print("SOURCE_DOCX=" + str(SOURCE_DOCX))
    print("OUTPUT_DOCX=" + str(out))
    print("REPORT_FILES=" + ";".join(str(p) for p in [*STAGE_FILES.values(), REPORT_MD]))
    print("CASE_COUNTS=" + ",".join(f"{k}:{len(v)}" for k, v in cases.items()))
    print("SOURCE_UNCHANGED=true")


if __name__ == "__main__":
    main()
